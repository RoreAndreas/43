"""Onglet Valorisation : DCF importé d'Excel et multiples réunis.

Le DCF de référence est celui de la note de valorisation qui a servi de
modèle : quatre exercices, g 2 %, flux en milieu d'année. Les classeurs sont
écrits par openpyxl — donc compressés comme ceux d'Excel — et importés par le
vrai champ de fichier. Les attendus sont recalculés ici en Python, au CMPC que
la page affiche, plutôt que recopiés.
"""

import openpyxl
import pytest

from conftest import mode_comparables

ONGLET = '.tabs button[data-tab="valorisation"]'

ANNEES = [2026, 2027, 2028, 2029]
EBIT = [54463, 63821, 76624, 106636]
DA = [63230, 80311, 100131, 82320]
CAPEX = [86641, 79152, 77427, 77427]
BFR = [-2723, -4404, -5212, -850]
G = 0.02
DETTE_NETTE = 129455


def classeur(chemin, *, annees=ANNEES, sans=(), capex_negatif=False, g=G, impot=0.25, ca=None):
    """Classeur au format de la note : libellés en A, exercices en en-tête."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DCF"
    ws.append(["M FCFA"] + [f"FY{str(a)[2:]}" for a in annees])
    if ca is not None:
        ws.append(["Chiffre d'affaires"] + ca[: len(annees)])
    lignes = {
        "EBIT (Résultat d'exploitation)": EBIT,
        "(+) Dotations nettes (D&A)": DA,
        "(−) Investissements (CapEx)": [-c for c in CAPEX] if capex_negatif else CAPEX,
        "(−) Variation du BFR": BFR,
    }
    for libelle, valeurs in lignes.items():
        if not any(m in libelle for m in sans):
            ws.append([libelle] + valeurs[: len(annees)])
    # Une ligne de calcul que l'import ne doit pas prendre pour l'EBIT.
    ws.append(["(−) Impôt normatif sur l'EBIT"] + [-e * 0.25 for e in EBIT[: len(annees)]])
    ws.append([])
    ws.append(["Croissance à l'infini (g)", g])
    ws.append(["Dette nette", DETTE_NETTE])
    if impot is not None:
        ws.append(["Taux d'IS", impot])
    wb.save(chemin)
    return chemin


def ve_attendue(wacc, g=G, impot=0.25):
    fcff = [e - max(0, e) * impot + d - c + b for e, d, c, b in zip(EBIT, DA, CAPEX, BFR)]
    facteurs = [(1 + wacc) ** -(t + 0.5) for t in range(len(fcff))]
    somme = sum(f * a for f, a in zip(fcff, facteurs))
    return somme + fcff[-1] * (1 + g) / (wacc - g) * facteurs[-1]


def montant(texte):
    t = texte.strip().replace(" ", "").replace(" ", "")
    signe = -1 if t.startswith("(") else 1
    return signe * float(t.strip("()"))


# Les quatre cases au-dessus du graphique, et les montants du résultat.
CASES = {"actualise": "somme", "vt": "vt", "cmpc": "cmpc", "tcam": "tcam"}


def carte(page, cle):
    if cle in ("ve", "vfp"):
        return page.inner_text(f'#valoDcf [data-o="{cle}"]')
    return page.inner_text(f'#valoDcf [data-case="{CASES[cle]}"] .dcf-case-valeur')


def importer(page, chemin):
    page.click(ONGLET)
    page.set_input_files("#valoFichier", str(chemin))
    page.wait_for_selector('#valoDcf [data-o="ve"], #valoDcf .valo-erreurs')


def cmpc(page):
    return page.evaluate("cmpcAuto()")


def test_sans_fichier_la_vue_invite_a_importer(page):
    page.click(ONGLET)
    assert page.is_visible('#valoDcf [data-action="importer"]')
    assert page.is_visible('#valoDcf [data-action="modele"]')
    assert page.query_selector("#valoDcf .dcf-hyp") is None
    assert page.query_selector('.tabs button[data-tab="comparables"]') is None


def test_le_dcf_importe_suit_la_note(page, tmp_path):
    importer(page, classeur(tmp_path / "dcf.xlsx"))
    ve = montant(carte(page, "ve"))
    assert abs(ve - ve_attendue(cmpc(page))) < 1
    assert [c.get_attribute("data-case") for c in page.query_selector_all("#valoDcf .dcf-case")] == [
        "somme", "vt", "cmpc", "tcam"]
    assert carte(page, "tcam").replace(" ", " ") == "25,1 %"
    assert abs(montant(carte(page, "vfp")) - (montant(carte(page, "ve")) - DETTE_NETTE)) <= 1
    # Le graphique : une colonne par exercice, deux barres chacune.
    colonnes = page.query_selector_all("#dcfVue .dcf-col")
    assert len(colonnes) == 4
    assert [a.inner_text() for a in page.query_selector_all("#dcfVue .dcf-axe-an")] == ["FY26", "FY27", "FY28", "FY29"]
    # Un clic bascule sur les lignes du calcul, un second revient au graphique.
    page.click("#dcfVue")
    detail = page.inner_text("#dcfVue .dcf-calc")
    assert "NOPAT" in detail and "40 847" in detail and "Facteur" in detail
    assert page.get_attribute("#dcfVue", "aria-pressed") == "true"
    page.click("#dcfVue")
    assert page.query_selector("#dcfVue .dcf-calc") is None
    assert len(page.query_selector_all("#dcfVue .dcf-col")) == 4


def test_la_note_est_retrouvee_a_son_cmpc(page, tmp_path):
    """À 12,02 %, le CMPC de la note, on retrouve sa VE de 739 340."""
    importer(page, classeur(tmp_path / "dcf.xlsx"))
    page.evaluate("model.brut.waccLocal = 0.1202; renderDcf()")
    assert abs(montant(carte(page, "ve")) - 739340) < 60
    assert abs(montant(carte(page, "actualise")) - 164452) < 10
    assert abs(montant(carte(page, "vt")) - 574888) < 60


def test_triangle_au_dela_de_65_pourcent(page, tmp_path):
    importer(page, classeur(tmp_path / "haut.xlsx"))
    assert page.query_selector('#valoDcf [data-case="vt"] .alerte-vt') is not None
    assert page.query_selector('#valoDcf .dcf-badge.is-alerte') is not None
    # À g très bas, la valeur terminale pèse moins.
    importer(page, classeur(tmp_path / "bas.xlsx", g=-0.30))
    assert page.query_selector('#valoDcf [data-case="vt"] .alerte-vt') is None


def test_le_resultat_mene_de_la_ve_a_la_vfp(page, tmp_path):
    """La sensibilité est retirée : le cadre de résultat porte seul le passage
    de la VE à la valeur des fonds propres, et la composition de la VE."""
    importer(page, classeur(tmp_path / "dcf.xlsx"))
    assert page.query_selector("#valoDcf .sens-table") is None
    assert page.query_selector("#valoDcf .pont-table") is None
    pont = [s.inner_text() for s in page.query_selector_all("#valoDcf .dcf-pont > span b")]
    ve, dette, vfp = (montant(x) for x in pont)
    assert ve == montant(carte(page, "ve"))
    assert dette == -DETTE_NETTE
    assert abs(ve + dette - vfp) <= 1
    assert vfp == montant(carte(page, "vfp"))
    # Le cadre de résultat descend jusqu'au bas du graphique.
    bas_resultat = page.eval_on_selector("#valoDcf .dcf-resultat", "e => e.getBoundingClientRect().bottom")
    bas_graphe = page.eval_on_selector("#dcfVue", "e => e.getBoundingClientRect().bottom")
    assert abs(bas_resultat - bas_graphe) <= 1


def test_facteur_a_deux_decimales(page, tmp_path):
    importer(page, classeur(tmp_path / "dcf.xlsx"))
    page.click("#dcfVue")
    ligne = [tr for tr in page.query_selector_all("#dcfVue .dcf-calc tbody tr")
             if "Facteur" in tr.inner_text()][0]
    for td in ligne.query_selector_all("td")[1:]:
        assert len(td.inner_text().split(",")[1]) == 2


def test_marge_d_ebitda_moyenne(page, tmp_path):
    """Moyenne des marges annuelles (EBIT + D&A) / CA ; sans CA, un tiret."""
    ca = [400000, 450000, 500000, 560000]
    importer(page, classeur(tmp_path / "ca.xlsx", ca=ca))
    attendu = sum((e + d) / c for e, d, c in zip(EBIT, DA, ca)) / len(ca)
    texte = page.inner_text("#valoDcf .dcf-hyp")
    assert "Dette nette" not in texte
    assert "Marge d'EBITDA moyenne" in texte
    assert f"{attendu * 100:.1f}".replace(".", ",") in texte
    importer(page, classeur(tmp_path / "sans.xlsx"))
    assert "CA absent du fichier" in page.inner_text("#valoDcf .dcf-hyp")


def test_capex_negatif_et_taux_d_is_du_pays(page, tmp_path):
    """Le CapEx se lit en montant, quel que soit son signe ; sans taux d'IS,
    celui du pays — 25 % par repli — est retenu."""
    importer(page, classeur(tmp_path / "dcf.xlsx", capex_negatif=True, impot=None))
    impot = page.evaluate("impotAuto()")
    assert abs(montant(carte(page, "ve")) - ve_attendue(cmpc(page), impot=impot)) < 1


def test_un_classeur_incomplet_est_refuse(page, tmp_path):
    importer(page, classeur(tmp_path / "trou.xlsx", sans=("CapEx",)))
    erreurs = page.inner_text("#valoDcf .valo-erreurs")
    assert "CapEx" in erreurs
    assert page.query_selector('#valoDcf [data-o="ve"]') is None


def test_le_modele_se_telecharge_et_se_reimporte(page, tmp_path):
    page.click(ONGLET)
    with page.expect_download() as dl:
        page.click('#valoDcf [data-action="modele"]')
    chemin = tmp_path / "modele.xlsx"
    dl.value.save_as(chemin)
    wb = openpyxl.load_workbook(chemin)
    ws = wb["DCF"]
    assert "Mode d'emploi" in wb.sheetnames
    annees = [c.value for c in ws[2][1:]]
    assert len(annees) == 5
    assert ws["A3"].value.startswith("Chiffre d'affaires")
    for i, valeurs in enumerate([EBIT, DA, CAPEX, BFR]):
        for j, v in enumerate(valeurs + [valeurs[-1]]):
            ws.cell(row=4 + i, column=2 + j, value=v)
    ws["B9"] = 0.02
    ws["B10"] = DETTE_NETTE
    wb.save(chemin)
    importer(page, chemin)
    assert page.query_selector('#valoDcf [data-o="ve"]') is not None
    assert "FY" + str(annees[-1])[2:] in page.inner_text("#valoDcf .valo-import")


def test_la_vue_comparables_montre_les_multiples(page):
    page.click(ONGLET)
    page.click('.valo-vues button[data-vue="comparables"]')
    assert page.is_visible("#comparablesAbsents")
    page.click('.tabs button[data-tab="params"]')
    mode_comparables(page)
    page.click(ONGLET)
    assert page.is_hidden("#comparablesAbsents")
    assert "VE / EBITDA" in page.inner_text("#comparablesMain")


def test_le_boulon_pose_les_donnees_de_test_puis_ramene_a_l_import(page):
    page.click(ONGLET)
    boulon = '#valoDcf .valo-boulon'
    # Il dépasse bien de la case : son centre est sur le bord haut de la zone d'import.
    cadre = page.eval_on_selector("#valoDcf .valo-depot", "e => e.getBoundingClientRect().top")
    haut = page.eval_on_selector(boulon, "e => e.getBoundingClientRect().top")
    bas = page.eval_on_selector(boulon, "e => e.getBoundingClientRect().bottom")
    assert haut < cadre < bas
    assert page.get_attribute(boulon, "aria-pressed") == "false"

    page.click(boulon)
    assert page.get_attribute(boulon, "aria-pressed") == "true"
    assert "Données de test" in page.inner_text("#valoDcf .valo-import")
    assert abs(montant(carte(page, "ve")) - ve_attendue(cmpc(page))) < 1
    # Ce sont les chiffres de la note : g = 2 %, dette nette 129 455, IS 25 %.
    pont = page.query_selector_all("#valoDcf .dcf-pont > span b")
    assert montant(pont[1].inner_text()) == -DETTE_NETTE

    page.click(boulon)
    assert page.is_visible('#valoDcf [data-action="importer"]')
    assert page.query_selector('#valoDcf [data-o="ve"]') is None


def test_le_boulon_ne_perd_pas_le_fichier_importe(page, tmp_path):
    importer(page, classeur(tmp_path / "mien.xlsx", g=0.01))
    mien = carte(page, "ve")
    page.click("#valoDcf .valo-boulon")
    assert carte(page, "ve") != mien
    page.click("#valoDcf .valo-boulon")
    assert carte(page, "ve") == mien
    assert "mien.xlsx" in page.inner_text("#valoDcf .valo-import")
    # Importer un fichier sort du mode test.
    page.click("#valoDcf .valo-boulon")
    page.set_input_files("#valoFichier", str(classeur(tmp_path / "autre.xlsx")))
    page.wait_for_function("document.querySelector('#valoDcf .valo-import').innerText.includes('autre.xlsx')")
    assert page.get_attribute("#valoDcf .valo-boulon", "aria-pressed") == "false"
